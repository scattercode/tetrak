#!/usr/bin/env python3
"""Generate the OCR comparison deck in the Tetrak OCR design language.

The deck walks the research: why archival material is hard, what engines exist,
five documents from the corpus, the results, a recommendation, and what is
still open. It is the same argument as /research/walkthrough/ on the site, in a
form you can stand up and talk through.

Nothing in it is typed in by hand twice:

  - The prose and structure come from `site/data/walkthrough.toml`, which the
    site page renders from as well, so the two cannot tell different stories.
  - Every figure comes from `evaluation/ocr/benchmark.csv` at generation time. The
    deck used to carry sentences like "0.91 against Claude's 0.98" written into
    the source, and they were wrong within a fortnight of the harness re-running.
  - The palette and faces come from `tools/tetrak_theme.py`, the one port of
    the site's design tokens into Python.

The output is a build artefact and is not committed: several megabytes of binary
that changes whenever the benchmark does. `collateral/` is gitignored.

Usage:
    python tools/generate_presentation.py
    python tools/generate_presentation.py --output collateral/my-deck.pptx
    python tools/generate_presentation.py --input-tokens 1800 --output-tokens 900

Requires `python-pptx` and `pdf2image` (+ poppler), both in the `docs` extra:
    pip install -e '.[docs]'
"""

from __future__ import annotations

import argparse
import csv
import tomllib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import tetrak_theme

ROOT = Path(__file__).resolve().parent.parent
EVAL_DIR = ROOT / "evaluation" / "ocr"
CORPUS_DIR = EVAL_DIR / "corpus" / "images"
WALKTHROUGH_TOML = ROOT / "site" / "data" / "walkthrough.toml"
DEFAULT_OUTPUT = ROOT / "collateral" / "ocr-tool-comparison-themed.pptx"

# PowerPoint reads neither PDF nor WebP, and its TIFF support varies by build,
# so every corpus item is rendered to PNG once and cached here. Gitignored.
CACHE_DIR = Path(__file__).resolve().parent / ".image-cache"


# ---------------------------------------------------------------------------
# Design
#
# Warm paper, one clay accent, Playfair Display over Source Serif 4, and no
# boxes -- hierarchy comes from the type scale and from whitespace. The values
# live in tetrak_theme, which is the port of site/assets/scss/_tokens.scss;
# nothing here should hard-code a colour or a face.
#
# This deck previously carried scattercode.dev's tokens instead: cobalt and
# yellow, Memphis pop-cards with hard offset shadows, Fraunces over Space
# Grotesk. If any of that reappears, it has come from an old copy of this file.
# ---------------------------------------------------------------------------
def rgb(name: str, palette: dict[str, str] | None = None) -> RGBColor:
    """A design token as a PowerPoint colour."""
    return RGBColor.from_string(tetrak_theme.token(name, palette).lstrip("#"))


PAPER = rgb("paper")
PAPER_2 = rgb("paper-2")
PAPER_3 = rgb("paper-3")
INK = rgb("ink")
INK_SOFT = rgb("ink-soft")
RULE = rgb("rule")
CLAY = rgb("clay")
CLAY_DEEP = rgb("clay-deep")

# Title and section slides invert. They use the site's own dark tokens rather
# than a black of their own, so a section break reads as Tetrak in dark mode
# rather than as a slide from somewhere else.
DARK = tetrak_theme.DARK
DARK_PAPER = rgb("paper", DARK)
DARK_INK = rgb("ink", DARK)
DARK_INK_SOFT = rgb("ink-soft", DARK)
DARK_CLAY = rgb("clay", DARK)

FONT_DISPLAY = tetrak_theme.FONT_DISPLAY
FONT_BODY = tetrak_theme.FONT_BODY
FONT_MONO = tetrak_theme.FONT_MONO

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.72)
CONTENT_W = SLIDE_W - 2 * MARGIN

# The backend that generated the ground truth. It is a ceiling reference, not a
# competitor: it is set back everywhere it appears and excluded from every
# "best local" comparison, exactly as the site's benchmark table does it.
CEILING = "claude"

# ---------------------------------------------------------------------------
# Anthropic API pricing, $ per million tokens.
# SOURCE: https://platform.claude.com/docs/en/about-claude/pricing
# VERIFIED: 2026-08-15. Pricing moves -- an earlier version of this deck quoted
# ~$0.035/page for Opus, which had already gone stale by the time anyone
# noticed. Re-verify against the URL above before quoting these anywhere, and
# update PRICING_VERIFIED_DATE when you do.
# ---------------------------------------------------------------------------
PRICING = {
    "Haiku 4.5": {"in": 1.00, "out": 5.00, "note": "budget model"},
    "Opus 4.8": {"in": 5.00, "out": 25.00, "note": "highest quality"},
}
PRICING_VERIFIED_DATE = "2026-08-15"
BATCH_DISCOUNT = 0.5
DEFAULT_INPUT_TOKENS_PER_PAGE = 1500
DEFAULT_OUTPUT_TOKENS_PER_PAGE = 750


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def new_slide(prs: Presentation, bg=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    font=FONT_BODY,
    size=16,
    bold=False,
    italic=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.35,
    tracking=0.0,
    wrap=True,
):
    """A text box, one paragraph per line.

    `tracking` is letter-spacing in points. PowerPoint stores it as the `spc`
    attribute on the run properties in hundredths of a point, which python-pptx
    does not surface, so it is set on the underlying element. The kickers need
    it -- mono uppercase without tracking reads as code, not as a label.
    """
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        if tracking:
            r.font._rPr.set("spc", str(int(round(tracking * 100))))
    return box


def add_rule(slide, x, y, w, color=RULE, weight=Inches(0.011)):
    """A hairline.

    The only structural mark in the system. Sections are separated by
    whitespace first and a rule second; nothing is ever put in a box.
    """
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, weight)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_panel(slide, x, y, w, h, fill=PAPER_2):
    """An inset surface: a fill and nothing else.

    No border, no shadow, no radius. This replaced `add_pop_card`, which drew
    the Memphis hard-offset shadow the previous design used on four slides.
    """
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_kicker(slide, x, y, text, color=INK_SOFT, width=Inches(6)):
    """The small mono label above a heading -- the site's `.kicker`."""
    return add_text(
        slide,
        x,
        y,
        width,
        Inches(0.26),
        text.upper(),
        font=FONT_MONO,
        size=10,
        color=color,
        tracking=1.2,
        line_spacing=1.0,
        wrap=False,
    )


def add_heading(slide, x, y, w, text, *, size=34, color=INK, h=Inches(0.9)):
    return add_text(
        slide,
        x,
        y,
        w,
        h,
        text,
        font=FONT_DISPLAY,
        size=size,
        bold=True,
        color=color,
        line_spacing=1.08,
    )


def add_numeral(slide, x, y, text, *, size=40, color=CLAY, w=Inches(0.9)):
    """A chapter numeral, set large in the margin.

    The previous design put these in filled circles. Tetrak has no disc and no
    second accent; the numeral carries itself at size.
    """
    return add_text(
        slide,
        x,
        y,
        w,
        Inches(0.8),
        text,
        font=FONT_DISPLAY,
        size=size,
        bold=True,
        color=color,
        line_spacing=1.0,
        wrap=False,
    )


def add_footer(slide, index, total, meta):
    label = meta["footer"]
    text = f"{label}  ·  {index} / {total}" if index is not None else f"{label}  ·  Appendix"
    add_text(
        slide,
        MARGIN,
        SLIDE_H - Inches(0.5),
        CONTENT_W,
        Inches(0.3),
        text,
        font=FONT_MONO,
        size=9,
        color=INK_SOFT,
        tracking=0.5,
        wrap=False,
    )


def add_image(slide, path: Path, x, y, w, h, *, frame=True):
    """Fit a picture inside a box, centred, with an optional hairline frame.

    The frame is `--rule` at hairline weight, which is what the site does to
    its own figures (`.hero-figure img`). It reads as a print frame around a
    photograph rather than as a UI box, which is why it is the one edge in the
    deck.
    """
    if not path or not path.exists():
        return None
    pic = slide.shapes.add_picture(str(path), x, y, height=h)
    if pic.width > w:
        scale = w / pic.width
        pic.width = int(w)
        pic.height = int(pic.height * scale)
    pic.left = int(x + (w - pic.width) / 2)
    pic.top = int(y + (h - pic.height) / 2)
    if frame:
        edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pic.left, pic.top, pic.width, pic.height)
        edge.fill.background()
        edge.line.color.rgb = RULE
        edge.line.width = Pt(0.75)
        edge.shadow.inherit = False
    return pic


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
def load_walkthrough() -> dict:
    with WALKTHROUGH_TOML.open("rb") as fh:
        return tomllib.load(fh)


def load_benchmark() -> tuple[dict[str, dict[str, str]], list[str]]:
    """Rows keyed by fixture, and the backends the harness actually wrote.

    Backends are discovered from the header rather than listed here, so adding
    one to the harness puts it in the deck without a code change -- the same
    contract the site's {{< benchmark >}} shortcode has.
    """
    path = EVAL_DIR / "benchmark.csv"
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        header = list(reader.fieldnames or [])
        rows = {row["fixture"]: row for row in reader}
    backends = [name[:-4] for name in header if name.endswith("_chr")]
    return rows, backends


def score(row, backend) -> tuple[float, float] | None:
    """Character similarity and word recall, or None where the engine cannot
    read the file at all."""
    chr_raw = row.get(f"{backend}_chr")
    wrd_raw = row.get(f"{backend}_wrd")
    if chr_raw in (None, "", "N/A") or wrd_raw in (None, "", "N/A"):
        return None
    return float(chr_raw), float(wrd_raw)


def best_local_score(row, backends) -> float | None:
    """The highest character similarity among the non-ceiling backends."""
    values = [pair[0] for b in backends if b != CEILING for pair in [score(row, b)] if pair]
    return max(values) if values else None


def is_best_local(row, backend, best) -> bool:
    """Whether this backend holds the best local score on this row.

    Every backend that matches is marked, not just the first one found. Ties
    are the normal case rather than an edge: auto-local does not compute a
    transcript of its own, it selects one of the other engines', so on most
    rows it scores exactly what the winning engine scored. Marking only one of
    them would make routing look like it had lost a race it had just won -- and
    the site's benchmark table marks every match, so the two would disagree.
    """
    if backend == CEILING or best is None:
        return False
    pair = score(row, backend)
    return pair is not None and pair[0] == best


def best_local_engine(row, engines) -> dict | None:
    """The engine to *name* as best local, where only one name will fit.

    Where a tie is with auto-local, auto-local is the answer: it is the thing
    you would actually run, and the engine it happened to pick is an
    implementation detail of that run.
    """
    best = best_local_score(row, [e["key"] for e in engines])
    matches = [e for e in engines if is_best_local(row, e["key"], best)]
    if not matches:
        return None
    for engine in matches:
        if engine["key"] == "auto-local":
            return engine
    return matches[0]


def engine_order(walkthrough, backends) -> list[str]:
    """Backends in the order the walkthrough introduces them.

    Anything the harness measured that the walkthrough has no prose for is
    appended rather than dropped: a new backend should show up in the results
    even before somebody has written a paragraph about it.
    """
    described = [e["key"] for e in walkthrough["engine"]]
    missing = [key for key in described if key not in backends]
    if missing:
        raise SystemExit(
            f"walkthrough.toml describes {missing} but evaluation/ocr/benchmark.csv has no "
            f"such column. Either the backend was renamed or the benchmark needs re-running."
        )
    return described + [b for b in backends if b not in described]


def engine_meta(walkthrough) -> dict[str, dict]:
    return {e["key"]: e for e in walkthrough["engine"]}


def accent_for(engine: dict) -> RGBColor:
    return rgb(engine.get("accent", "ink-soft"))


def render_source(filename: str, max_px: int = 1600) -> Path | None:
    """Render a corpus item to a cached PNG.

    The corpus is mixed: JPEG, PNG, a 21-page PDF and two archival TIFFs.
    PowerPoint will not open the PDF, and its TIFF handling varies between
    builds, so everything is normalised here rather than trusting the reader.
    """
    src = CORPUS_DIR / filename
    if not src.exists():
        return None
    out = CACHE_DIR / f"{Path(filename).stem}.png"
    if out.exists():
        return out

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    if src.suffix.lower() == ".pdf":
        from pdf2image import convert_from_path

        image = convert_from_path(str(src), dpi=150, first_page=1, last_page=1)[0]
    else:
        from PIL import Image

        image = Image.open(src)
        # Archival TIFFs arrive bilevel or palettised; PowerPoint wants neither.
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")

    if image.width > max_px:
        from PIL import Image as PILImage

        height = round(image.height * max_px / image.width)
        image = image.resize((max_px, height), PILImage.LANCZOS)
    image.save(out, "PNG")
    return out


def fmt_pair(pair: tuple[float, float] | None) -> str:
    return "N/A" if pair is None else f"{pair[0]:.2f} / {pair[1]:.2f}"


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def slide_title(prs, meta):
    s = new_slide(prs, bg=DARK_PAPER)
    add_rule(s, MARGIN, Inches(1.32), Inches(2.2), color=DARK_CLAY, weight=Inches(0.02))
    add_kicker(s, MARGIN, Inches(0.92), "Research walkthrough", color=DARK_CLAY)
    # Two lines of it, deliberately. The title is set at display size and the
    # slide is laid out around a headline that wraps -- at a size that fits on
    # one line it stops carrying the slide. Anything appreciably longer than
    # the current title will need the size dropping rather than this loosening.
    add_text(
        s,
        MARGIN,
        Inches(1.78),
        CONTENT_W,
        Inches(2.2),
        meta["title"].upper(),
        font=FONT_DISPLAY,
        size=54,
        bold=True,
        color=DARK_INK,
        line_spacing=1.02,
    )
    add_text(
        s,
        MARGIN,
        Inches(3.95),
        Inches(10),
        Inches(0.6),
        meta["subtitle"],
        font=FONT_BODY,
        size=21,
        italic=True,
        color=DARK_CLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(4.8),
        Inches(7.6),
        Inches(1.4),
        meta["standfirst"],
        font=FONT_BODY,
        size=13,
        color=DARK_INK_SOFT,
        line_spacing=1.45,
    )
    add_text(
        s,
        MARGIN,
        Inches(6.35),
        Inches(8),
        Inches(0.3),
        meta["byline"],
        font=FONT_MONO,
        size=11,
        color=DARK_INK,
        wrap=False,
    )
    add_text(
        s,
        MARGIN,
        Inches(6.75),
        Inches(8),
        Inches(0.3),
        f"{meta['footer']}  ·  {meta['date']}",
        font=FONT_MONO,
        size=10,
        color=DARK_INK_SOFT,
        tracking=0.5,
        wrap=False,
    )
    return s


def slide_section(prs, meta, title, subtitle, index, total):
    s = new_slide(prs, bg=DARK_PAPER)
    add_rule(s, MARGIN, Inches(2.6), Inches(2.2), color=DARK_CLAY, weight=Inches(0.02))
    add_text(
        s,
        MARGIN,
        Inches(3.0),
        Inches(11),
        Inches(1.1),
        title.upper(),
        font=FONT_DISPLAY,
        size=42,
        bold=True,
        color=DARK_INK,
        line_spacing=1.0,
    )
    if subtitle:
        add_text(
            s,
            MARGIN,
            Inches(4.1),
            Inches(10),
            Inches(0.6),
            subtitle,
            font=FONT_BODY,
            size=16,
            italic=True,
            color=DARK_CLAY,
        )
    add_footer(s, index, total, meta)
    return s


def slide_contents(prs, meta, chapters, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "Contents")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What we'll cover")

    y = Inches(2.05)
    row_h = Inches(0.95)
    for i, chapter in enumerate(chapters):
        if i:
            add_rule(s, MARGIN, y - Inches(0.12), CONTENT_W)
        add_numeral(s, MARGIN, y - Inches(0.08), chapter["number"], size=32)
        add_text(
            s,
            MARGIN + Inches(0.85),
            y,
            Inches(3.1),
            row_h,
            chapter["title"],
            font=FONT_DISPLAY,
            size=19,
            bold=True,
            color=INK,
        )
        add_text(
            s,
            MARGIN + Inches(4.1),
            y + Inches(0.04),
            CONTENT_W - Inches(4.1),
            row_h,
            chapter["blurb"],
            size=13,
            color=INK_SOFT,
        )
        y += row_h
    add_footer(s, index, total, meta)
    return s


def slide_challenge(prs, meta, challenges, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "The challenge")
    add_heading(
        s,
        MARGIN,
        Inches(0.95),
        Inches(9.5),
        "Thousands of documents await digitisation,\nand every one of them is different",
        size=30,
        h=Inches(1.6),
    )

    gap = Inches(0.42)
    w = (CONTENT_W - 3 * gap) / 4
    y = Inches(3.3)
    for i, item in enumerate(challenges):
        x = MARGIN + i * (w + gap)
        add_rule(s, x, y, w, color=CLAY, weight=Inches(0.022))
        add_numeral(s, x, y + Inches(0.28), str(i + 1), size=30)
        add_text(
            s,
            x,
            y + Inches(1.0),
            w,
            Inches(0.8),
            item["title"],
            font=FONT_DISPLAY,
            size=17,
            bold=True,
            color=INK,
            line_spacing=1.12,
        )
        add_text(
            s,
            x,
            y + Inches(1.85),
            w,
            Inches(1.9),
            item["body"],
            size=12,
            color=INK_SOFT,
            line_spacing=1.42,
        )
    add_footer(s, index, total, meta)
    return s


def slide_engines(prs, meta, engines, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "The engines")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What we measured, and what each one is")

    cols = 4
    gap = Inches(0.4)
    w = (CONTENT_W - (cols - 1) * gap) / cols
    row_h = Inches(2.55)
    y0 = Inches(2.0)
    for i, engine in enumerate(engines):
        r, c = divmod(i, cols)
        x = MARGIN + c * (w + gap)
        y = y0 + r * (row_h + Inches(0.2))
        add_rule(s, x, y, w, color=accent_for(engine), weight=Inches(0.022))
        add_text(
            s,
            x,
            y + Inches(0.2),
            w,
            Inches(0.24),
            engine["availability"].upper(),
            font=FONT_MONO,
            size=8.5,
            color=INK_SOFT,
            tracking=1.0,
            line_spacing=1.0,
        )
        add_text(
            s,
            x,
            y + Inches(0.52),
            w,
            Inches(0.4),
            engine["label"],
            font=FONT_DISPLAY,
            size=18,
            bold=True,
            color=INK,
            line_spacing=1.0,
        )
        add_text(
            s,
            x,
            y + Inches(1.0),
            w,
            row_h - Inches(1.05),
            engine["body"],
            size=11.5,
            color=INK_SOFT,
            line_spacing=1.42,
        )
    add_footer(s, index, total, meta)
    return s


def slide_document(prs, meta, doc, position, hero_count, benchmark, engines, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "The corpus")
    add_text(
        s,
        MARGIN,
        Inches(0.98),
        Inches(6),
        Inches(0.3),
        f"Document {position} of {hero_count}",
        font=FONT_MONO,
        size=11,
        color=CLAY_DEEP,
        wrap=False,
    )
    add_heading(s, MARGIN, Inches(1.4), Inches(6.6), doc["title"], size=27, h=Inches(1.2))
    add_text(
        s,
        MARGIN,
        Inches(2.62),
        Inches(6.6),
        Inches(0.35),
        doc["subtitle"],
        size=13,
        italic=True,
        color=INK_SOFT,
    )
    add_rule(s, MARGIN, Inches(3.12), Inches(6.6), color=CLAY, weight=Inches(0.022))
    add_text(
        s,
        MARGIN,
        Inches(3.36),
        Inches(6.6),
        Inches(1.55),
        doc["challenge"],
        size=13.5,
        color=INK,
        line_spacing=1.5,
    )

    row = benchmark.get(doc["file"])
    if row:
        add_scoreboard(s, MARGIN, Inches(5.05), Inches(6.6), row, engines)

    add_image(s, render_source(doc["file"]), Inches(8.0), Inches(1.15), Inches(4.6), Inches(5.5))
    add_footer(s, index, total, meta)
    return s


def add_scoreboard(slide, x, y, w, row, engines):
    """The per-engine figures for one document, read from the CSV.

    Two mono columns, best local in bold and the ceiling set back -- the same
    marking the site's benchmark table uses, so a reader who has seen one
    recognises the other.
    """
    best = best_local_score(row, [e["key"] for e in engines])
    add_text(
        slide,
        x,
        y,
        w,
        Inches(0.24),
        "Character similarity / word recall",
        font=FONT_MONO,
        size=8.5,
        color=INK_SOFT,
        tracking=1.0,
        line_spacing=1.0,
        wrap=False,
    )

    # Two columns. Seven engines stacked in one ran off the bottom of the
    # slide, under the footer; splitting them keeps the block inside the plate
    # and level with the foot of the scan beside it.
    line_h = Inches(0.235)
    per_col = -(-len(engines) // 2)
    col_w = w / 2
    for i, engine in enumerate(engines):
        col, rowi = divmod(i, per_col)
        cx = x + col * col_w
        ty = y + Inches(0.34) + rowi * line_h
        key = engine["key"]
        is_ceiling = key == CEILING
        is_best = is_best_local(row, key, best)
        color = INK_SOFT if is_ceiling else INK
        label = f"{engine['label']}*" if is_ceiling else engine["label"]
        add_text(
            slide,
            cx,
            ty,
            Inches(1.55),
            line_h,
            label,
            font=FONT_MONO,
            size=10,
            bold=is_best,
            color=color,
            line_spacing=1.0,
            wrap=False,
        )
        add_text(
            slide,
            cx + Inches(1.6),
            ty,
            Inches(1.4),
            line_h,
            fmt_pair(score(row, key)),
            font=FONT_MONO,
            size=10,
            bold=is_best,
            color=color,
            line_spacing=1.0,
            wrap=False,
        )
    add_text(
        slide,
        x,
        y + Inches(0.34) + per_col * line_h + Inches(0.1),
        w,
        Inches(0.24),
        "* ceiling reference, not a measurement · bold is best local",
        font=FONT_MONO,
        size=8,
        italic=True,
        color=INK_SOFT,
        line_spacing=1.0,
        wrap=False,
    )


def slide_contact_sheet(prs, meta, documents, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "The corpus")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "All nine, on one theme")
    add_text(
        s,
        MARGIN,
        Inches(1.72),
        Inches(10.5),
        Inches(0.6),
        "Los Angeles stage and picture-palace ephemera, c. 1910–1945 — chosen to span the "
        "failure modes that matter in archive digitisation rather than to flatter anything.",
        size=12.5,
        color=INK_SOFT,
    )

    top_row, bottom_row = documents[:5], documents[5:]
    y = Inches(2.35)
    for items in (top_row, bottom_row):
        if not items:
            continue
        gap = Inches(0.3)
        w = (CONTENT_W - (len(items) - 1) * gap) / len(items)
        for i, doc in enumerate(items):
            x = MARGIN + i * (w + gap)
            add_image(s, render_source(doc["file"], max_px=900), x, y, w, Inches(1.85))
            add_text(
                s,
                x,
                y + Inches(1.95),
                w,
                Inches(0.3),
                doc["file"],
                font=FONT_MONO,
                size=7.5,
                color=INK_SOFT,
                line_spacing=1.15,
            )
        y += Inches(2.5)
    add_footer(s, index, total, meta)
    return s


def slide_matrix(prs, meta, benchmark, documents, engines, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.55), "The results")
    add_heading(s, MARGIN, Inches(0.86), CONTENT_W, "Every engine against every document", size=27)
    add_text(
        s,
        MARGIN,
        Inches(1.5),
        CONTENT_W,
        Inches(0.34),
        "Character similarity / word recall, scored against the reference transcripts.",
        size=12,
        color=INK_SOFT,
    )

    headers = ["Document"] + [
        f"{e['label']}*" if e["key"] == CEILING else e["label"] for e in engines
    ]
    body_rows = [d for d in documents if d["file"] in benchmark]
    has_average = "Average" in benchmark
    rows = len(body_rows) + 1 + (1 if has_average else 0)

    x, y = MARGIN, Inches(1.95)
    w, h = CONTENT_W, Inches(4.55)
    table = s.shapes.add_table(rows, len(headers), x, y, w, h).table
    # PowerPoint's default table style bands rows in the theme's accent. Off:
    # the marking here is bold-for-best, and banding fights it.
    table.first_row = False
    table.horz_banding = False
    table.columns[0].width = Inches(2.85)
    for c in range(1, len(headers)):
        table.columns[c].width = int((w - Inches(2.85)) / (len(headers) - 1))

    def style_cell(
        cell,
        text,
        *,
        font=FONT_MONO,
        size=10,
        bold=False,
        color=INK,
        fill=PAPER,
        align=PP_ALIGN.CENTER,
    ):
        cell.text = text
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Pt(5)
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        p = cell.text_frame.paragraphs[0]
        p.alignment = align
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color

    for c, text in enumerate(headers):
        style_cell(
            table.cell(0, c),
            text,
            font=FONT_BODY,
            size=11,
            bold=True,
            color=PAPER,
            fill=INK,
            align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER,
        )

    for r, doc in enumerate(body_rows, start=1):
        row = benchmark[doc["file"]]
        best = best_local_score(row, [e["key"] for e in engines])
        ground = PAPER if r % 2 else PAPER_2
        style_cell(
            table.cell(r, 0),
            doc["title"],
            font=FONT_BODY,
            size=10.5,
            color=INK,
            fill=ground,
            align=PP_ALIGN.LEFT,
        )
        for c, engine in enumerate(engines, start=1):
            key = engine["key"]
            is_ceiling = key == CEILING
            style_cell(
                table.cell(r, c),
                fmt_pair(score(row, key)),
                size=9.5,
                bold=is_best_local(row, key, best),
                color=INK_SOFT if is_ceiling else INK,
                fill=ground,
            )

    if has_average:
        avg = benchmark["Average"]
        last = rows - 1
        style_cell(
            table.cell(last, 0),
            "Average",
            font=FONT_BODY,
            size=11,
            bold=True,
            color=INK,
            fill=PAPER_3,
            align=PP_ALIGN.LEFT,
        )
        best = best_local_score(avg, [e["key"] for e in engines])
        for c, engine in enumerate(engines, start=1):
            key = engine["key"]
            style_cell(
                table.cell(last, c),
                fmt_pair(score(avg, key)),
                size=9.5,
                bold=is_best_local(avg, key, best),
                color=INK_SOFT if key == CEILING else INK,
                fill=PAPER_3,
            )

    add_text(
        s,
        MARGIN,
        Inches(6.65),
        CONTENT_W,
        Inches(0.6),
        "Bold marks the best local engine per document. * claude generated the ground truth "
        "and is a ceiling reference rather than a score. N/A means the engine cannot read that "
        "file at all. vision is macOS-only and is not in this run.",
        size=10,
        italic=True,
        color=INK_SOFT,
        line_spacing=1.35,
    )
    add_footer(s, index, total, meta)
    return s


def slide_chart(prs, meta, benchmark, engines, title, kicker, note, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), kicker)
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, title, size=27)

    avg = benchmark["Average"]
    labels, values, colors = [], [], []
    for engine in engines:
        pair = score(avg, engine["key"])
        labels.append(engine["label"])
        values.append(pair[0] if pair else 0.0)
        colors.append(accent_for(engine))

    data = CategoryChartData()
    data.categories = labels
    data.add_series("Average character similarity", values)
    frame = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.8), CONTENT_W, Inches(4.4), data
    )
    chart = frame.chart
    chart.has_legend = False
    chart.font.name = FONT_MONO
    chart.font.size = Pt(11)
    chart.font.color.rgb = INK_SOFT

    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0.00"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(11)
    plot.data_labels.font.name = FONT_MONO
    plot.data_labels.font.color.rgb = INK

    series = plot.series[0]
    for i, color in enumerate(colors):
        point = series.points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color

    chart.value_axis.maximum_scale = 1.0
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RULE
    chart.value_axis.format.line.color.rgb = RULE
    chart.category_axis.format.line.color.rgb = RULE

    if note:
        add_text(
            s,
            MARGIN,
            Inches(6.4),
            CONTENT_W,
            Inches(0.7),
            note,
            size=11.5,
            color=INK_SOFT,
            line_spacing=1.4,
        )
    add_footer(s, index, total, meta)
    return s


def slide_recommendation(prs, meta, recommendations, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "The results")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "Which engine for which material?")

    y = Inches(2.05)
    row_h = Inches(0.72)
    add_text(
        s,
        MARGIN,
        y - Inches(0.34),
        Inches(6),
        Inches(0.26),
        "Your material",
        font=FONT_MONO,
        size=9,
        color=INK_SOFT,
        tracking=1.0,
        wrap=False,
    )
    add_text(
        s,
        MARGIN + Inches(5.5),
        y - Inches(0.34),
        Inches(3),
        Inches(0.26),
        "Use",
        font=FONT_MONO,
        size=9,
        color=INK_SOFT,
        tracking=1.0,
        wrap=False,
    )
    add_text(
        s,
        MARGIN + Inches(7.6),
        y - Inches(0.34),
        Inches(4),
        Inches(0.26),
        "Why",
        font=FONT_MONO,
        size=9,
        color=INK_SOFT,
        tracking=1.0,
        wrap=False,
    )

    for item in recommendations:
        add_rule(s, MARGIN, y, CONTENT_W)
        add_text(
            s,
            MARGIN,
            y + Inches(0.16),
            Inches(5.3),
            Inches(0.5),
            item["material"],
            size=13,
            color=INK,
        )
        engine = item["engine"]
        if item.get("second"):
            engine = f"{engine}   ·   then {item['second']}"
        add_text(
            s,
            MARGIN + Inches(5.5),
            y + Inches(0.19),
            Inches(2.2),
            Inches(0.5),
            engine,
            font=FONT_MONO,
            size=11,
            color=CLAY_DEEP,
            wrap=False,
        )
        add_text(
            s,
            MARGIN + Inches(7.6),
            y + Inches(0.19),
            CONTENT_W - Inches(7.6),
            Inches(0.5),
            item["why"],
            size=11.5,
            color=INK_SOFT,
        )
        y += row_h
    add_rule(s, MARGIN, y, CONTENT_W)
    add_footer(s, index, total, meta)
    return s


def slide_findings(prs, meta, findings, benchmark, engines, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "Summary")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What the evidence says")

    y = Inches(2.05)
    for i, finding in enumerate(findings):
        add_numeral(s, MARGIN, y - Inches(0.06), str(i + 1), size=30)
        add_text(
            s,
            MARGIN + Inches(0.85),
            y,
            CONTENT_W - Inches(0.85),
            Inches(0.4),
            finding["title"],
            font=FONT_DISPLAY,
            size=18,
            bold=True,
            color=INK,
            line_spacing=1.1,
        )
        add_text(
            s,
            MARGIN + Inches(0.85),
            y + Inches(0.44),
            CONTENT_W - Inches(0.85),
            Inches(0.85),
            finding["body"],
            size=12.5,
            color=INK_SOFT,
            line_spacing=1.42,
        )
        y += Inches(1.42)

    # The headline numbers, so the claims above have their evidence on the same
    # slide. Read from the CSV like everything else.
    avg = benchmark["Average"]
    keys = [e["key"] for e in engines]
    best = best_local_engine(avg, engines)
    meta_by_key = {e["key"]: e for e in engines}
    panel_y = y + Inches(0.1)
    add_panel(s, MARGIN, panel_y, CONTENT_W, Inches(1.05))
    add_text(
        s,
        MARGIN + Inches(0.3),
        panel_y + Inches(0.18),
        Inches(6),
        Inches(0.24),
        "Corpus average, character similarity / word recall",
        font=FONT_MONO,
        size=8.5,
        color=INK_SOFT,
        tracking=1.0,
        wrap=False,
    )
    parts = []
    if best:
        parts.append(f"best local — {best['label']} {fmt_pair(score(avg, best['key']))}")
    if CEILING in keys:
        parts.append(f"ceiling — {meta_by_key[CEILING]['label']} {fmt_pair(score(avg, CEILING))}")
    add_text(
        s,
        MARGIN + Inches(0.3),
        panel_y + Inches(0.5),
        CONTENT_W - Inches(0.6),
        Inches(0.4),
        "     ".join(parts),
        font=FONT_MONO,
        size=13,
        color=INK,
        wrap=False,
    )
    add_footer(s, index, total, meta)
    return s


def slide_questions(prs, meta, questions, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "What follows")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What is still open")

    cols = 2
    gap = Inches(0.6)
    w = (CONTENT_W - gap) / cols
    y0 = Inches(2.0)
    row_h = Inches(1.62)
    for i, question in enumerate(questions):
        r, c = divmod(i, cols)
        x = MARGIN + c * (w + gap)
        y = y0 + r * row_h
        add_rule(s, x, y, w, color=CLAY, weight=Inches(0.022))
        add_text(
            s,
            x,
            y + Inches(0.2),
            w,
            Inches(0.5),
            question["title"],
            font=FONT_DISPLAY,
            size=16,
            bold=True,
            color=INK,
            line_spacing=1.15,
        )
        add_text(
            s,
            x,
            y + Inches(0.78),
            w,
            Inches(0.75),
            question["body"],
            size=11.5,
            color=INK_SOFT,
            line_spacing=1.4,
        )
    add_footer(s, index, total, meta)
    return s


def slide_pricing(prs, meta, input_tokens, output_tokens, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "Appendix")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What the cloud path would cost", size=27)
    add_text(
        s,
        MARGIN,
        Inches(1.62),
        CONTENT_W,
        Inches(0.35),
        f"Estimated per scanned page, assuming ~{input_tokens} input and ~{output_tokens} "
        "output tokens — roughly what this corpus runs at.",
        size=12.5,
        color=INK_SOFT,
    )

    def page_cost(model, batch=False):
        rate = PRICING[model]
        discount = BATCH_DISCOUNT if batch else 1.0
        return (
            input_tokens / 1_000_000 * rate["in"] + output_tokens / 1_000_000 * rate["out"]
        ) * discount

    entries = [
        ("Haiku 4.5", PRICING["Haiku 4.5"]["note"], page_cost("Haiku 4.5")),
        ("Opus 4.8", PRICING["Opus 4.8"]["note"], page_cost("Opus 4.8")),
        ("Opus 4.8, Batch API", "50% discount", page_cost("Opus 4.8", batch=True)),
    ]
    gap = Inches(0.5)
    w = (CONTENT_W - 2 * gap) / 3
    y = Inches(2.35)
    for i, (name, note, cost) in enumerate(entries):
        x = MARGIN + i * (w + gap)
        add_rule(s, x, y, w, color=CLAY, weight=Inches(0.022))
        add_text(
            s,
            x,
            y + Inches(0.22),
            w,
            Inches(0.35),
            name,
            font=FONT_DISPLAY,
            size=18,
            bold=True,
            color=INK,
            line_spacing=1.0,
        )
        add_text(
            s,
            x,
            y + Inches(0.62),
            w,
            Inches(0.3),
            note,
            size=12,
            italic=True,
            color=INK_SOFT,
        )
        add_text(
            s,
            x,
            y + Inches(1.05),
            w,
            Inches(0.8),
            f"~${cost:.3f}",
            font=FONT_DISPLAY,
            size=40,
            bold=True,
            color=CLAY_DEEP,
            line_spacing=1.0,
        )
        add_text(
            s,
            x,
            y + Inches(1.85),
            w,
            Inches(0.3),
            "per page",
            size=12,
            color=INK,
        )
        add_text(
            s,
            x,
            y + Inches(2.2),
            w,
            Inches(0.3),
            f"~${cost * 1000:,.0f} per 1,000 pages",
            font=FONT_MONO,
            size=10,
            color=INK_SOFT,
        )

    low = min(page_cost(m, batch=True) for m in PRICING) * 50_000
    high = max(page_cost(m, batch=False) for m in PRICING) * 50_000
    add_text(
        s,
        MARGIN,
        Inches(5.35),
        CONTENT_W,
        Inches(1.2),
        "The local engines have no per-page cost but do have setup and staff time. A practical "
        "pattern is to run auto-local over everything and send only what lands in workspace/triage/ to the "
        f"API. At full volume a 50,000-page collection would run roughly ${low:,.0f}–${high:,.0f} "
        "depending on model — but the first question is whether the material may leave the "
        "building at all.",
        size=12.5,
        color=INK,
        line_spacing=1.45,
    )
    add_text(
        s,
        MARGIN,
        Inches(6.75),
        CONTENT_W,
        Inches(0.3),
        f"Pricing verified {PRICING_VERIFIED_DATE} against platform.claude.com — re-check "
        "before quoting these.",
        font=FONT_MONO,
        size=8.5,
        italic=True,
        color=INK_SOFT,
    )
    add_footer(s, index, total, meta)
    return s


def slide_engine_detail(prs, meta, engines, index, total):
    s = new_slide(prs)
    add_kicker(s, MARGIN, Inches(0.62), "Appendix")
    add_heading(s, MARGIN, Inches(0.95), CONTENT_W, "What each engine is best at", size=27)

    cols = 2
    gap = Inches(0.6)
    w = (CONTENT_W - gap) / cols
    y0 = Inches(1.9)
    row_h = Inches(1.32)
    for i, engine in enumerate(engines):
        r, c = divmod(i, cols)
        x = MARGIN + c * (w + gap)
        y = y0 + r * row_h
        add_rule(s, x, y, w, color=accent_for(engine))
        add_text(
            s,
            x,
            y + Inches(0.14),
            Inches(2.2),
            Inches(0.32),
            engine["label"],
            font=FONT_MONO,
            size=13,
            color=INK,
            wrap=False,
        )
        add_text(
            s,
            x + Inches(2.3),
            y + Inches(0.16),
            w - Inches(2.3),
            Inches(0.3),
            engine["availability"],
            font=FONT_MONO,
            size=9,
            color=INK_SOFT,
            wrap=False,
        )
        add_text(
            s,
            x,
            y + Inches(0.52),
            w,
            Inches(0.34),
            f"Best for — {engine['best_for']}",
            size=11.5,
            color=INK,
            line_spacing=1.3,
        )
        add_text(
            s,
            x,
            y + Inches(0.88),
            w,
            Inches(0.34),
            f"Limits — {engine['limits']}",
            size=11.5,
            color=INK_SOFT,
            line_spacing=1.3,
        )
    add_footer(s, index, total, meta)
    return s


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def build(output: Path, input_tokens: int, output_tokens: int) -> Path:
    walkthrough = load_walkthrough()
    meta = walkthrough["meta"]
    benchmark, backends = load_benchmark()

    by_key = engine_meta(walkthrough)
    ordered = engine_order(walkthrough, backends)
    engines = [
        by_key.get(key, {"key": key, "label": key, "availability": "", "body": ""})
        for key in ordered
    ]
    local_engines = [e for e in engines if e["key"] != CEILING]

    documents = walkthrough["document"]
    heroes = [d for d in documents if d.get("hero")]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total = 15
    slide_title(prs, meta)
    slide_contents(prs, meta, walkthrough["chapter"], 2, total)
    slide_challenge(prs, meta, walkthrough["challenge"], 3, total)
    slide_engines(prs, meta, engines, 4, total)
    for n, doc in enumerate(heroes, start=1):
        slide_document(prs, meta, doc, n, len(heroes), benchmark, engines, 4 + n, total)
    slide_contact_sheet(prs, meta, documents, 10, total)
    slide_matrix(prs, meta, benchmark, documents, engines, 11, total)
    slide_chart(
        prs,
        meta,
        benchmark,
        local_engines,
        "Choosing per file beats choosing an engine",
        "The results",
        "auto-local outperforms every individual local engine by running the ones eligible for "
        "each file and keeping the best-scoring transcript. Averages over the whole corpus; the "
        "ceiling reference is excluded.",
        12,
        total,
    )
    slide_recommendation(prs, meta, walkthrough["recommendation"], 13, total)
    slide_findings(prs, meta, walkthrough["finding"], benchmark, engines, 14, total)
    slide_questions(prs, meta, walkthrough["question"], 15, total)

    slide_section(
        prs, meta, "Appendix", "Cost, engine detail, and the full comparison", None, total
    )
    slide_pricing(prs, meta, input_tokens, output_tokens, None, total)
    slide_engine_detail(prs, meta, engines, None, total)
    slide_chart(
        prs,
        meta,
        benchmark,
        engines,
        "Every engine, including the ceiling",
        "Appendix",
        "EasyOCR and PaddleOCR cannot read PDFs, so their averages exclude the souvenir "
        "programme. claude generated the ground truth and is a ceiling reference, not a "
        "competitor.",
        None,
        total,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output .pptx path")
    parser.add_argument(
        "--input-tokens",
        type=int,
        default=DEFAULT_INPUT_TOKENS_PER_PAGE,
        help="Assumed input tokens/page for the cost slide",
    )
    parser.add_argument(
        "--output-tokens",
        type=int,
        default=DEFAULT_OUTPUT_TOKENS_PER_PAGE,
        help="Assumed output tokens/page for the cost slide",
    )
    args = parser.parse_args()
    out = build(args.output, args.input_tokens, args.output_tokens)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
